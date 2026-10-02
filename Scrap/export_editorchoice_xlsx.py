"""Exporta a coleta existente, recuperando autores e imagens do snapshot local."""
import json
import hashlib
import re
from pathlib import Path
from urllib.parse import urljoin, urlparse

from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parent


def main():
    records = json.loads((BASE / 'trading_view_editorchoice.json').read_text(encoding='utf-8'))
    snapshot = BASE / 'base de indicadores editor choice v3_20260904_235525.htm'
    cards = {}
    for page in re.split(r'<!-- ===== PAGE \d+ ===== -->', snapshot.read_text(encoding='utf-8')):
        soup = BeautifulSoup(page, 'lxml')
        for card in soup.find_all('article'):
            title = card.find('a', attrs={'data-qa-id': 'ui-lib-card-link-title'})
            if title:
                cards[urlparse(title.get('href')).path.split('/')[2].split('-')[0]] = card
    image_dir = BASE / 'base de indicadores  editor choice v2_files'
    local_count = 0
    for record in records:
        card = cards.get(urlparse(record['url']).path.split('/')[2].split('-')[0])
        if card is None:
            cache = Path('D:/Documentos/Projetos/TradingView-scrap/tv_cache_html') / ('inner_' + hashlib.sha1(record['url'].encode()).hexdigest() + '.html')
            card = BeautifulSoup(cache.read_text(encoding='utf-8'), 'lxml')
        author = card.select_one('address a[href]') or card.select_one('a[href^="/u/"]')
        if author:
            record['url_autor'] = urljoin('https://www.tradingview.com', author['href'])
            record['autor'] = urlparse(record['url_autor']).path.strip('/').split('/')[-1]
        preview = card.select_one('a[data-qa-id="ui-lib-card-link-image"] img')
        record['url_imagem'] = preview.get('src', '') if preview else ''
        if not record['url_imagem']:
            meta = card.find('meta', property='og:image')
            record['url_imagem'] = meta.get('content', '') if meta else ''
        filename = Path(urlparse(record['url_imagem']).path).name
        local = image_dir / filename
        record['imagem_local'] = str(local) if filename and local.is_file() else ''
        local_count += bool(record['imagem_local'])
    assert len(records) == len({r['url'] for r in records})
    assert all(r['autor'] and r['url_autor'] and r['url_imagem'] for r in records)
    columns = [('posicao', 'Posição'), ('titulo', 'Indicador'), ('url', 'Link do indicador'),
               ('autor', 'Autor'), ('url_autor', 'Link do autor'), ('data_atualizacao', 'Atualização'),
               ('boosts', 'Boosts'), ('comentarios', 'Comentários'), ('views', 'Visualizações'),
               ('tags', 'Tags'), ('descricao', 'Descrição'), ('url_imagem', 'Link da imagem'),
               ('imagem_local', 'Imagem local')]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Editor Choice'
    ws.append([label for _, label in columns])
    for record in records:
        ws.append([record.get(key, '') for key, _ in columns])
        for col, (key, _) in enumerate(columns, 1):
            cell = ws.cell(ws.max_row, col)
            cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if key in ('posicao', 'boosts', 'comentarios', 'views') and str(cell.value).isdigit():
                cell.value = int(cell.value)
            target = record['url'] if key == 'titulo' else record.get(key, '')
            if target and key in ('titulo', 'url', 'url_autor', 'url_imagem', 'imagem_local'):
                cell.hyperlink = Path(target).as_uri() if key == 'imagem_local' else target
                cell.style = 'Hyperlink'
        ws.row_dimensions[ws.max_row].height = 48
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17365D')
    for column, width in zip('ABCDEFGHIJKLM', [10, 48, 42, 24, 38, 16, 12, 14, 18, 48, 80, 42, 48]):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = 'C2'
    table = Table(displayName='IndicadoresEditorChoice', ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    notes = wb.create_sheet('Sobre a base')
    for row in [
        ['Item', 'Informação'], ['Indicadores', len(records)],
        ['Fonte', snapshot.name], ['Coleta', '04–05/09/2026; exportação sem nova captura'],
        ['Autores', 'Recuperados do HTML local'],
        ['Imagens locais correspondentes', local_count],
        ['Pasta de imagens existentes', str(image_dir)],
        ['Imagens', 'Links de prévia disponíveis para todos; imagens não incorporadas ao XLSX.'],
        ['Descrição', 'Resumo da coleta original, limitado a aproximadamente 300 caracteres.'],
    ]:
        notes.append(row)
    notes.column_dimensions['A'].width = 36
    notes.column_dimensions['B'].width = 110
    output = BASE / 'trading_view_editorchoice.xlsx'
    wb.save(output)
    check = load_workbook(output)
    sheet = check['Editor Choice']
    assert sheet.max_row == len(records) + 1
    for row in sheet.iter_rows(min_row=2):
        assert all(row[i].hyperlink and row[i].hyperlink.target for i in (1, 2, 4, 11))
    print(json.dumps({'xlsx': str(output), 'records': len(records), 'authors': sum(bool(r['autor']) for r in records),
                      'image_links': sum(bool(r['url_imagem']) for r in records), 'local_images': local_count}, ensure_ascii=True))


if __name__ == '__main__':
    main()
